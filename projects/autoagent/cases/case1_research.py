#!/usr/bin/env python3
"""Кейс 1 — Research-ассистент: Triage/Planner → Web Surfer + File Surfer → отчёт.

Флоу:
  1. Planner разбивает вопрос на 2-4 подзадачи + поисковые запросы.
  2. Web Surfer ищет в интернете (DuckDuckGo) и вытаскивает текст страниц.
  3. File Surfer ищет ответы в локальной папке sources/ (md/txt/pdf).
  4. Compiler собирает отчёт со ссылками [n] в markdown (+docx).

Запуск:
  python case1_research.py                         # вопрос по умолчанию
  python case1_research.py "ваш вопрос"
  python case1_research.py "вопрос" --pages 8
"""

import os, re, sys, json, time, shutil
from datetime import datetime
from urllib.parse import urlparse, parse_qs, unquote

import httpx
from bs4 import BeautifulSoup

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(BASE))
from auto_zen import llm, safe, ROLE_MODELS  # noqa: E402

OUT_DIR = os.path.join(BASE, "out")
SRC_DIR = os.path.join(BASE, "sources")
os.makedirs(OUT_DIR, exist_ok=True)


def make_run_dir(prefix: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(OUT_DIR, f"{prefix}_{ts}")
    os.makedirs(os.path.join(path, "sources"), exist_ok=True)
    return path

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36"}

DEFAULT_QUESTION = "Как оценивать качество RAG-систем в продакшене?"


# ─── Planner ─────────────────────────────────────────────────

def plan_question(question: str) -> list[dict]:
    prompt = (
        "Ты — Planner исследовательского агента. Разбей вопрос на 2-4 подзадачи. "
        "Для каждой подзадачи дай короткий поисковый запрос на английском.\n"
        "Формат ответа — строго по одной строке на подзадачу:\n"
        "- <подзадача на русском> | <english search query>\n"
        "Никаких других строк."
    )
    raw = llm(prompt, question, temp=0.2, role="planner")
    tasks = []
    for line in raw.splitlines():
        line = line.strip()
        if not line.startswith("-"):
            continue
        body = line.lstrip("- ").strip()
        if "|" in body:
            sub, query = body.split("|", 1)
        else:
            sub, query = body, body
        sub, query = sub.strip(), query.strip()
        if sub and query:
            tasks.append({"subtask": sub, "query": query})
    return tasks[:4] or [{"subtask": question, "query": question}]


# ─── Web Surfer ──────────────────────────────────────────────

def ddg_search(query: str, limit: int = 5) -> list[dict]:
    try:
        r = httpx.get(
            "https://html.duckduckgo.com/html/",
            params={"q": query},
            headers=UA,
            timeout=30,
            follow_redirects=True,
        )
        r.raise_for_status()
    except Exception as e:
        print(safe(f"    [web] поиск не удался: {e}"))
        return []
    soup = BeautifulSoup(r.text, "html.parser")
    results = []
    for block in soup.select(".result"):
        a = block.select_one(".result__a")
        if not a:
            continue
        href = a.get("href", "")
        if "uddg=" in href:
            try:
                href = unquote(parse_qs(urlparse(href).query)["uddg"][0])
            except Exception:
                pass
        if href.startswith("//"):
            href = "https:" + href
        snippet_el = block.select_one(".result__snippet")
        results.append({
            "title": a.get_text(" ", strip=True),
            "url": href,
            "snippet": snippet_el.get_text(" ", strip=True) if snippet_el else "",
        })
        if len(results) >= limit:
            break
    return results


def fetch_page(url: str, max_chars: int = 20000) -> str:
    try:
        r = httpx.get(url, headers=UA, timeout=25, follow_redirects=True)
        if r.status_code >= 400 or "text/html" not in r.headers.get("content-type", ""):
            return ""
    except Exception:
        return ""
    soup = BeautifulSoup(r.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "aside", "form", "noscript"]):
        tag.decompose()
    text = re.sub(r"\s+", " ", soup.get_text(" ", strip=True))
    return text[:max_chars]


def web_surfer(tasks: list[dict], max_pages: int) -> list[dict]:
    sources, seen_urls, seen_domains = [], set(), set()
    for task in tasks:
        if len(sources) >= max_pages:
            break
        print(safe(f"    [web] «{task['query']}»"))
        for res in ddg_search(task["query"], limit=4):
            if len(sources) >= max_pages:
                break
            url = res["url"]
            domain = urlparse(url).netloc
            if not url.startswith("http") or url in seen_urls or domain in seen_domains:
                continue
            text = fetch_page(url)
            if len(text) < 200:
                continue
            seen_urls.add(url)
            seen_domains.add(domain)
            sources.append({
                "title": res["title"],
                "url": url,
                "text": text,
                "subtask": task["subtask"],
            })
            print(safe(f"      + {domain} ({len(text)} симв.)"))
            time.sleep(0.4)
    return sources


# ─── File Surfer ─────────────────────────────────────────────

def read_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".md", ".txt"):
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext == ".pdf":
        try:
            from pdfminer.high_level import extract_text
            return extract_text(path)
        except Exception as e:
            print(safe(f"    [docs] PDF не прочитан ({os.path.basename(path)}): {e}"))
    return ""


def file_surfer(question: str, tasks: list[dict], limit: int = 2) -> list[dict]:
    if not os.path.isdir(SRC_DIR):
        return []
    words = {w.lower() for w in re.findall(r"\w{4,}", question + " " + " ".join(t["subtask"] for t in tasks))}
    scored = []
    for name in sorted(os.listdir(SRC_DIR)):
        path = os.path.join(SRC_DIR, name)
        if not os.path.isfile(path):
            continue
        text = read_document(path)
        if not text.strip():
            continue
        doc_words = {w.lower() for w in re.findall(r"\w{4,}", text)}
        score = len(words & doc_words)
        scored.append((score, name, text, path))
    scored.sort(key=lambda x: (-x[0], x[1]))
    docs = []
    for score, name, text, path in scored[:limit]:
        docs.append({"title": f"Локальный документ: {name}", "url": "", "text": text[:3000],
                     "subtask": "локальные материалы", "local_path": path})
        print(safe(f"    [docs] {name} (совпадений: {score})"))
    return docs


def materialize(run_dir: str, sources: list[dict]) -> None:
    src_dir = os.path.join(run_dir, "sources")
    os.makedirs(src_dir, exist_ok=True)
    for i, s in enumerate(sources, 1):
        if s.get("local_path"):
            name = f"{i:02d}_{os.path.basename(s['local_path'])}"
            shutil.copy2(s["local_path"], os.path.join(src_dir, name))
        else:
            slug = re.sub(r"[^\w\-]+", "_", s["title"], flags=re.UNICODE)[:60].strip("_") or f"source_{i}"
            name = f"{i:02d}_{slug}.md"
            with open(os.path.join(src_dir, name), "w", encoding="utf-8") as f:
                f.write(f"# {s['title']}\n\nURL: {s['url']}\n\n{s['text']}")
        s["file"] = f"sources/{name}"


# ─── Compiler ────────────────────────────────────────────────

def compile_report(question: str, tasks: list[dict], sources: list[dict]) -> str:
    blocks = []
    for i, s in enumerate(sources, 1):
        head = f"[{i}] {s['title']}" + (f" — {s['url']}" if s["url"] else "")
        blocks.append(f"{head}\nВыдержка: {s['text'][:3500]}")
    src_text = "\n\n".join(blocks)
    tasks_text = "\n".join(f"- {t['subtask']}" for t in tasks)
    prompt = (
        "Ты — Compiler исследовательского агента. Напиши отчёт на русском в markdown "
        "строго по предоставленным источникам.\n"
        "Требования:\n"
        "1. Каждое фактическое утверждение сопровождай ссылкой [n] на источник.\n"
        "2. Если источники расходятся — укажи расхождение.\n"
        "3. Если данных нет — напиши 'в источниках не найдено', не выдумывай.\n"
        "4. Структура: краткий ответ (3-5 пунктов), разделы по подзадачам, "
        "проверка источников, ограничения, список источников.\n"
    )
    user = (
        f"Вопрос: {question}\n\nПодзадачи:\n{tasks_text}\n\n"
        f"Источники:\n{src_text}"
    )
    report = llm(prompt, user, temp=0.2, role="reporter")
    used = sorted({int(n) for n in re.findall(r"\[(\d+)\]", report)})
    print(safe(f"    [compiler] использовано ссылок: {used or 'нет'}"))
    if not used:
        report += "\n\n> Внимание: модель не проставила ссылки — проверьте вывод вручную.\n"
    return report


def md_to_docx(md: str, path: str) -> bool:
    try:
        from docx import Document
    except ImportError:
        return False
    doc = Document()
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip() or line.strip() in ("---", "***"):
            continue
        stripped = line.replace("**", "").replace("`", "")
        if line.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", stripped), style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+\.\s+", "", stripped), style="List Number")
        elif line.startswith(">"):
            style = doc.styles["Quote"] if "Quote" in [s.name for s in doc.styles] else None
            doc.add_paragraph(stripped.lstrip("> ").strip(), style=style) if style else doc.add_paragraph(stripped.lstrip("> ").strip())
        else:
            doc.add_paragraph(stripped)
    doc.save(path)
    return True


# ─── Main ────────────────────────────────────────────────────

def main():
    argv = sys.argv[1:]
    demo = "--demo" in argv
    if demo:
        argv.remove("--demo")
    max_pages = 6
    if "--pages" in argv:
        i = argv.index("--pages")
        try:
            max_pages = int(argv[i + 1])
        except Exception:
            pass
        del argv[i:i + 2]
    question = " ".join(argv).strip() or DEFAULT_QUESTION

    run_dir = make_run_dir("research")
    print(safe(f"\n{'='*60}\n  Кейс 1. Research-ассистент\n  Вопрос: {question}\n{'='*60}"))
    if not demo:
        print(safe(f"  Папка прогона: {run_dir}"))
    print(safe("\n[1/4] Planner: декомпозиция вопроса..."))
    tasks = plan_question(question)
    for t in tasks:
        print(safe(f"    - {t['subtask']}  (query: {t['query']})"))

    print(safe("\n[2/4] Web Surfer: поиск и загрузка страниц..."))
    web = web_surfer(tasks, max_pages)

    print(safe("\n[3/4] File Surfer: локальные документы..."))
    local = file_surfer(question, tasks)
    sources = local + web

    if not sources:
        print(safe("  [ERROR] Не найдено ни одного источника."))
        return

    materialize(run_dir, sources)
    print(safe(f"\n  Источников собрано: {len(sources)} (локальных: {len(local)}, web: {len(web)})"))
    print(safe("\n[4/4] Compiler: отчёт со ссылками..."))
    report = compile_report(question, tasks, sources)

    header = "" if report.lstrip().startswith("# ") else f"# Research-отчёт: {question}\n\n"
    md_path = os.path.join(run_dir, "report.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(header + report)

    docx_path = os.path.join(run_dir, "report.docx")
    md_to_docx(header + report, docx_path)

    info = {
        "question": question,
        "created": datetime.now().isoformat(timespec="seconds"),
        "models": ROLE_MODELS,
        "tasks": tasks,
        "sources": [{"id": i, "kind": "local" if s.get("local_path") else "web",
                     "title": s["title"], "url": s["url"], "file": s["file"]}
                    for i, s in enumerate(sources, 1)],
    }
    with open(os.path.join(run_dir, "sources.json"), "w", encoding="utf-8") as f:
        json.dump(info, f, ensure_ascii=False, indent=2)

    if demo:
        print(safe("\n  [Done] Готово: report.md · report.docx · sources.json · sources/"))
    else:
        print(safe(f"\n  [Done] ВСЁ в одной папке — открой её:\n  {run_dir}"))
        print(safe("    report.md · report.docx · sources.json · sources/ (полные тексты)"))


if __name__ == "__main__":
    main()
