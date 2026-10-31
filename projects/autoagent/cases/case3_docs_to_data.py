#!/usr/bin/env python3
"""Кейс 3 — Документы → данные: File Surfer → Coding Agent → CSV + сводка.

Флоу:
  1. File Surfer читает документы (txt/md/docx/pdf).
  2. Агент извлекает поля в строгий JSON (vendor, invoice_no, date, items, суммы).
  3. Coding Agent + pandas нормализует и валидирует данные (арифметика, дубликаты, даты).
  4. Результат: CSV, items CSV, сводка md и SQL для записи в PostgreSQL через MCP.

Запуск:
  python case3_docs_to_data.py
"""

import os, sys, re, json, shutil
from datetime import datetime
import pandas as pd

BASE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(BASE))
from auto_zen import llm, safe, ROLE_MODELS  # noqa: E402

DOCS_DIR = os.path.join(BASE, "sample_docs")
OUT_DIR = os.path.join(BASE, "out")
os.makedirs(OUT_DIR, exist_ok=True)


def make_run_dir(prefix: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = os.path.join(OUT_DIR, f"{prefix}_{ts}")
    os.makedirs(path, exist_ok=True)
    return path

SCHEMA = """{
  "vendor": "string",
  "invoice_no": "string",
  "date": "YYYY-MM-DD или null",
  "currency": "USD | EUR | ...",
  "items": [{"description": "string", "qty": число, "unit_price": число}],
  "subtotal": число или null,
  "tax": число или null,
  "total": число или null
}"""


# ─── File Surfer ────────────────────────────────────────────

def read_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts)
        except Exception as e:
            print(safe(f"    [docs] DOCX не прочитан ({os.path.basename(path)}): {e}"))
            return ""
    if ext == ".pdf":
        try:
            from pdfminer.high_level import extract_text
            return extract_text(path)
        except Exception as e:
            print(safe(f"    [docs] PDF не прочитан ({os.path.basename(path)}): {e}"))
            return ""
    return ""


# ─── Extraction ─────────────────────────────────────────────

def parse_json(raw: str):
    t = raw.strip()
    t = re.sub(r"^```(?:json)?\s*", "", t)
    t = re.sub(r"\s*```$", "", t)
    start, end = t.find("{"), t.rfind("}")
    if start == -1 or end <= start:
        return None
    try:
        return json.loads(t[start:end + 1])
    except Exception:
        return None


def extract_record(text: str, name: str, extra: str = "") -> dict | None:
    system = ("Ты — агент извлечения данных из документов. "
              "Верни ТОЛЬКО валидный JSON без markdown и пояснений.")
    user = (
        f"Документ «{name}»:\n{text}\n\n"
        f"Верни JSON строго по схеме:\n{SCHEMA}\n\n"
        "Правила: суммы — числами без символов валют и разделителей тысяч; "
        "дата — в формате YYYY-MM-DD; отсутствующее поле — null."
    )
    if extra:
        user += "\n\n" + extra
    data = parse_json(llm(system, user, temp=0.0, role="planner"))
    if data is None:
        data = parse_json(llm(system, user + "\n\nПредыдущий ответ не был валидным JSON. Верни только JSON.", temp=0.0, role="planner"))
    return data


def to_num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    s = re.sub(r"[^\d.,-]", "", s)
    if not s:
        return None
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def norm_date(v):
    if not v:
        return None
    dt = pd.to_datetime(str(v), errors="coerce")
    if pd.isna(dt):
        return None
    return dt.strftime("%Y-%m-%d")


# ─── Main ───────────────────────────────────────────────────

def main():
    run_dir = make_run_dir("docs_run")
    shutil.copytree(DOCS_DIR, os.path.join(run_dir, "documents"))
    print(safe(f"\n{'='*60}\n  Кейс 3. Документы -> данные\n{'='*60}"))
    print(safe(f"  Папка прогона: {run_dir}"))

    names = sorted(n for n in os.listdir(DOCS_DIR)
                   if os.path.splitext(n)[1].lower() in (".txt", ".md", ".docx", ".pdf"))
    print(safe(f"\n[1/4] File Surfer: найдено документов — {len(names)}"))

    rows, items_rows, raw_records = [], [], {}
    for name in names:
        print(safe(f"\n  [2/4] Извлечение полей: {name}"))
        text = read_document(os.path.join(DOCS_DIR, name))
        if not text.strip():
            rows.append({"file": name, "invoice_no": None, "vendor": None, "date": None,
                         "currency": None, "items_count": 0, "subtotal": None, "tax": None,
                         "total": None, "status": "CHECK", "issues": "не удалось прочитать документ"})
            continue
        rec = extract_record(text, name) or {}
        raw_records[name] = rec
        items = rec.get("items") or []
        items_norm = []
        for it in items:
            qty = to_num(it.get("qty"))
            price = to_num(it.get("unit_price"))
            items_norm.append({
                "description": str(it.get("description") or ""),
                "qty": qty, "unit_price": price,
                "amount": round(qty * price, 2) if qty is not None and price is not None else None,
            })
        row = {
            "file": name,
            "invoice_no": (str(rec.get("invoice_no")).strip() if rec.get("invoice_no") else None),
            "vendor": (str(rec.get("vendor")).strip() if rec.get("vendor") else None),
            "date": norm_date(rec.get("date")),
            "currency": (str(rec.get("currency")).upper() if rec.get("currency") else "USD"),
            "items_count": len(items_norm),
            "subtotal": to_num(rec.get("subtotal")),
            "tax": to_num(rec.get("tax")),
            "total": to_num(rec.get("total")),
        }
        print(safe(f"    vendor={row['vendor']} | no={row['invoice_no']} | "
                   f"date={row['date']} | total={row['total']} {row['currency']}"))
        rows.append(row)
        for it in items_norm:
            items_rows.append({"file": name, "invoice_no": row["invoice_no"], "vendor": row["vendor"], **it})

    print(safe(f"\n[3/4] Coding Agent + pandas: нормализация и валидация..."))
    df = pd.DataFrame(rows)
    df["issues"] = [[] for _ in range(len(df))]

    for idx, r in df.iterrows():
        miss = [label for key, label in (("vendor", "vendor"), ("invoice_no", "invoice_no"),
                                         ("date", "дата"), ("total", "total")) if not r.get(key)]
        if miss:
            df.at[idx, "issues"].append("нет поля: " + ", ".join(miss))
        if r["total"] is not None and r["total"] <= 0:
            df.at[idx, "issues"].append("total <= 0")
        if r["subtotal"] is not None and r["tax"] is not None and r["total"] is not None:
            if abs(r["subtotal"] + r["tax"] - r["total"]) > 0.01:
                df.at[idx, "issues"].append(
                    f"арифметика: {r['subtotal']:.2f} + {r['tax']:.2f} != {r['total']:.2f}")
        item_sum = sum(it["amount"] for it in items_rows
                       if it["file"] == r["file"] and it["amount"] is not None)
        if r["subtotal"] is not None and item_sum and abs(item_sum - r["subtotal"]) > 0.05:
            df.at[idx, "issues"].append(
                f"сумма позиций {item_sum:.2f} != subtotal {r['subtotal']:.2f}")

    keys = df["invoice_no"].fillna("").astype(str).str.strip().str.upper()
    dup_mask = keys.ne("") & keys.duplicated(keep=False)
    for idx in df.index[dup_mask]:
        df.at[idx, "issues"].append("дубликат invoice_no")

    df["issues"] = df["issues"].apply(lambda xs: "; ".join(xs))
    df["status"] = df["issues"].apply(lambda s: "OK" if not s else "CHECK")

    cols = ["file", "invoice_no", "vendor", "date", "currency", "items_count",
            "subtotal", "tax", "total", "status", "issues"]
    df = df[cols]

    csv_path = os.path.join(run_dir, "invoices.csv")
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    print(safe(f"    [Done] CSV: {csv_path}"))

    items_df = pd.DataFrame(items_rows, columns=["file", "invoice_no", "vendor", "description", "qty", "unit_price", "amount"])
    items_path = os.path.join(run_dir, "items.csv")
    items_df.to_csv(items_path, index=False, encoding="utf-8-sig")
    print(safe(f"    [Done] Items CSV: {items_path}"))

    json_path = os.path.join(run_dir, "extracted.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(raw_records, f, ensure_ascii=False, indent=2)
    print(safe(f"    [Done] Raw JSON: {json_path}"))

    ok_n = int((df["status"] == "OK").sum())
    bad_n = int((df["status"] != "OK").sum())
    valid_totals = df[df["status"] == "OK"].groupby("currency")["total"].sum().to_dict()

    print(safe("\n[4/4] Итоги:"))
    print(safe(f"    документов: {len(df)} | OK: {ok_n} | требуют проверки: {bad_n}"))
    for cur, total in valid_totals.items():
        print(safe(f"    сумма OK ({cur}): {total:,.2f}"))
    for _, r in df.iterrows():
        if r["status"] != "OK":
            print(safe(f"    [!] {r['file']}: {r['issues']}"))

    lines = ["# Сводка: документы → данные", "",
             f"- Документов: {len(df)}",
             f"- Без замечаний: {ok_n}",
             f"- Требуют проверки: {bad_n}"]
    for cur, total in valid_totals.items():
        lines.append(f"- Сумма по валидным ({cur}): {total:,.2f}")
    lines += ["", "| Файл | Invoice | Vendor | Дата | Total | Статус | Замечания |",
              "|---|---|---|---|---|---|---|"]
    for _, r in df.iterrows():
        vals = [r["file"], r["invoice_no"] or "", r["vendor"] or "", r["date"] or "",
                f"{r['total']:.2f} {r['currency']}" if r["total"] is not None else "",
                r["status"], r["issues"] or ""]
        lines.append("| " + " | ".join(str(v).replace("|", "\\|") for v in vals) + " |")
    lines += ["", "> Запись в PostgreSQL выполняется через MCP: готовый SQL — `insert.sql`."]
    summary_path = os.path.join(run_dir, "summary.md")
    with open(summary_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(safe(f"    [Done] Сводка: {summary_path}"))

    sql_lines = ["-- INSERT для PostgreSQL (выполнить через MCP-инструмент)",
                 "INSERT INTO invoices (invoice_no, vendor, invoice_date, currency, subtotal, tax, total)",
                 "VALUES"]
    vals = []
    def sql_quote_str(v):
        if v is None or (hasattr(pd, "isna") and pd.isna(v)):
            return "NULL"
        s = str(v).replace("\\", "\\\\").replace("'", "''")
        return f"'{s}'"

    def sql_quote_num(v):
        if v is None or (hasattr(pd, "isna") and pd.isna(v)):
            return "NULL"
        try:
            return str(float(v))
        except (ValueError, TypeError):
            return "NULL"

    for _, r in df[df["status"] == "OK"].iterrows():
        inv_no = sql_quote_str(r.get("invoice_no"))
        vendor = sql_quote_str(r.get("vendor"))
        inv_date = sql_quote_str(r.get("date"))
        currency = sql_quote_str(r.get("currency"))
        subtotal = sql_quote_num(r.get("subtotal"))
        tax = sql_quote_num(r.get("tax"))
        total = sql_quote_num(r.get("total"))
        vals.append(f"  ({inv_no}, {vendor}, {inv_date}, {currency}, {subtotal}, {tax}, {total})")
    if vals:
        sql_lines.append(",\n".join(vals))
        sql_lines.append("ON CONFLICT (invoice_no) DO NOTHING;")
    else:
        sql_lines.append("  -- нет валидных строк")
    sql_path = os.path.join(run_dir, "insert.sql")
    with open(sql_path, "w", encoding="utf-8") as f:
        f.write("\n".join(sql_lines))
    with open(os.path.join(run_dir, "run_info.json"), "w", encoding="utf-8") as f:
        json.dump({
            "created": datetime.now().isoformat(timespec="seconds"),
            "models": ROLE_MODELS,
            "documents": len(df),
            "ok": ok_n,
            "check": bad_n,
            "valid_totals": {k: round(v, 2) for k, v in valid_totals.items()},
        }, f, ensure_ascii=False, indent=2)

    print(safe(f"\n  [Done] ВСЁ в одной папке — открой её:\n  {run_dir}"))
    print(safe("    summary.md · invoices.csv · items.csv · extracted.json · insert.sql · documents/"))


if __name__ == "__main__":
    main()
